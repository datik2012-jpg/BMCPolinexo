"""Bounded, memory-only Har HaBituach import. Never log source values."""
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
from zipfile import ZipFile, BadZipFile

from defusedxml.ElementTree import iterparse
from openpyxl import load_workbook

MAX_EXPANDED = 25 * 1024 * 1024
MAX_ROWS = 10000
FIELDS = ('insured_id', 'category', 'subcategory', 'product_type', 'insurer', 'period',
          'additional_details', 'premium', 'frequency', 'policy_number', 'classification')
HEADERS = ('תעודת זהות', 'ענף ראשי', 'ענף (משני)', 'סוג מוצר', 'חברה', 'תקופת ביטוח',
           'פרטים נוספים', 'פרמיה בש"ח', 'סוג פרמיה', 'מספר פוליסה', 'סיווג תכנית')

class ImportProblem(ValueError):
    pass

def display(value):
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.strftime('%d/%m/%Y')
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()

def canonical(value):
    return re.sub(r'\s+', ' ', display(value)).replace('״', '"')

def normalize(fields):
    values = dict(fields)
    issues = []
    values['frequency'] = {'חודשית': 'monthly', 'חודשי': 'monthly', 'שנתית': 'annual', 'שנתי': 'annual'}.get(fields['frequency'], fields['frequency'])
    if values['frequency'] not in ('monthly', 'annual'):
        issues.append('תדירות הפרמיה חסרה או אינה מוכרת')
    try:
        premium_text = fields['premium'].strip().removeprefix('₪').strip()
        if not re.fullmatch(r'(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d{1,2})?', premium_text):
            raise InvalidOperation
        amount = Decimal(premium_text.replace(',', ''))
        if not amount.is_finite() or amount < 0 or amount > Decimal('999999999999') or amount.as_tuple().exponent < -2:
            raise InvalidOperation
        values['premium'] = format(amount, 'f')
    except (InvalidOperation, ValueError):
        issues.append('הפרמיה חסרה או אינה תקינה')
    if any(not values[key] for key in ('insured_id', 'insurer', 'policy_number')):
        issues.append('חסרים פרטים הנדרשים לשיוך לפוליסה')
    return values, issues

def inspect_archive(payload):
    try:
        with ZipFile(BytesIO(payload)) as archive:
            members = archive.infolist()
            if len(members) > 2000 or sum(m.file_size for m in members) > MAX_EXPANDED:
                raise ImportProblem('הקובץ גדול מדי לעיבוד')
            if len({m.filename for m in members}) != len(members):
                raise ImportProblem('מבנה הקובץ אינו תקין')
            for member in members:
                if member.flag_bits & 1 or member.filename.endswith('vbaProject.bin'):
                    raise ImportProblem('הקובץ אינו נתמך')
                if member.filename.startswith('xl/worksheets/') and member.filename.endswith('.xml'):
                    with archive.open(member) as stream:
                        rows = 0
                        for _, element in iterparse(stream, events=('end',)):
                            tag = element.tag.rsplit('}', 1)[-1]
                            if tag == 'row':
                                rows += 1
                                if rows > MAX_ROWS or int(element.attrib.get('r', rows)) > MAX_ROWS:
                                    raise ImportProblem('הקובץ מכיל יותר מדי שורות')
                            if tag == 'c':
                                address = element.attrib.get('r', '')
                                letters = re.sub('[0-9]', '', address)
                                column = 0
                                for letter in letters:
                                    column = column * 26 + ord(letter) - 64
                                if column > 64:
                                    raise ImportProblem('מבנה העמודות אינו נתמך')
                            element.clear()
    except ImportProblem:
        raise
    except Exception:
        raise ImportProblem('לא ניתן לקרוא את קובץ ה־Excel') from None

def parse_workbook(payload):
    inspect_archive(payload)
    workbook = None
    try:
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False, keep_links=False)
        all_entries = []
        report_date = None
        data_sheets = 0
        for sheet in workbook.worksheets:
            sheet.reset_dimensions()
            columns = None
            entries = []
            for number, row in enumerate(sheet.iter_rows(max_col=64), 1):
                texts = [display(cell.value) for cell in row]
                if not any(texts):
                    continue
                normalized = [canonical(t) for t in texts]
                if all(h in normalized for h in HEADERS):
                    columns = {field: normalized.index(header) for field, header in zip(FIELDS, HEADERS)}
                    continue
                if columns is None:
                    for cell in texts:
                        if re.fullmatch(r'\d{2}/\d{2}/\d{4}', cell):
                            report_date = cell
                    continue
                original = {field: texts[index] for field, index in columns.items()}
                identifier_warning = False
                for field in ('insured_id', 'policy_number'):
                    cell = row[columns[field]]
                    if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        if re.fullmatch(r'0+', cell.number_format or ''):
                            original[field] = original[field].zfill(len(cell.number_format))
                        else:
                            identifier_warning = True
                # A section heading has no coverage-level content.
                if original['category'].startswith('תחום') and not any(original[k] for k in FIELDS if k != 'category'):
                    continue
                if not any(original.values()):
                    continue
                if any(row[index].data_type == 'f' for index in columns.values()):
                    raise ImportProblem('הקובץ מכיל נוסחאות. יש להעלות ייצוא עם ערכים בלבד')
                values, issues = normalize(original)
                if identifier_warning:
                    issues.append('מזהה נשמר כמספר ב־Excel; יש לבדוק אם חסרים אפסים מובילים')
                entries.append({'id': f'{len(all_entries) + len(entries) + 1}', 'source_sheet': sheet.title,
                                'source_row': number, 'original': original, 'values': values, 'issues': issues})
            if entries:
                data_sheets += 1
                all_entries.extend(entries)
        if data_sheets > 1:
            raise ImportProblem('יש להעלות קובץ עם גיליון נתונים אחד בלבד')
        if not all_entries:
            raise ImportProblem('לא נמצאו נתוני ביטוח במבנה הנתמך')
        insured = {e['values']['insured_id'] for e in all_entries if e['values']['insured_id']}
        if len(insured) > 1:
            raise ImportProblem('הקובץ מכיל יותר ממבוטח אחד. יש להעלות קובץ נפרד לכל מבוטח')
        counts = Counter(tuple(e['values'][f] for f in FIELDS) for e in all_entries)
        for entry in all_entries:
            if counts[tuple(entry['values'][f] for f in FIELDS)] > 1:
                entry['issues'].append('חשד לשורה כפולה — יש לבדוק לפני החרגה')
        warnings = ['יש נתונים הדורשים בדיקה לפני השימוש בתוצאות'] if any(e['issues'] for e in all_entries) else []
        return {'report_date': report_date, 'entries': all_entries, 'warnings': warnings}
    except ImportProblem:
        raise
    except Exception:
        raise ImportProblem('לא ניתן לקרוא את קובץ ה־Excel') from None
    finally:
        if workbook is not None:
            workbook.close()
