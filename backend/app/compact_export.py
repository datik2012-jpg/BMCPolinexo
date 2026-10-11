"""Compact RTL presentation with native Excel selectors and auditable inputs."""
from collections import Counter
from datetime import date, datetime
from itertools import zip_longest
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from math import ceil
import re
import json
import unicodedata
from functools import lru_cache
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter

from .parser import FIELDS, HEADERS, normalize

NAVY, TEAL, PALE, BLUE, LINE = '003366', '0066CC', 'EAF4FA', '65BCE3', 'C6DDEB'
MONEY = '#,##0.00 "₪"'
FAMILY = 'כל המשפחה'
# Shared physical columns align customer and policy tables with the reference.
WIDTHS = [14, 20, 16, 17, 12, 15, 14, 10, 16, 18, 11, 13, 20]
SPANS = [(1, 1), (2, 2), (3, 3), (4, 5), (6, 6), (7, 8), (9, 9), (10, 10), (11, 13)]
TITLES = ['', 'שם חברת ביטוח', 'מספר פוליסה', 'כיסויים', 'תאריך תחילת ביטוח', 'מבוטחים', 'פרמיה חודשית', 'סיווג תכנית', 'הוראות הסוכן']


def parsed_date(value):
    for pattern in ('%Y-%m-%d', '%d/%m/%Y', '%d.%m.%Y'):
        try:
            return datetime.strptime(value.strip(), pattern).date()
        except ValueError:
            pass
    return None


def insurance_age(birth_date, on_date):
    born = parsed_date(birth_date)
    if not born or born > on_date:
        return None
    return on_date.year - born.year - ((on_date.month, on_date.day) < (born.month, born.day))


def start_date(period):
    # Recognize a complete leading date; never infer one from a year or label.
    match = re.match(r'^\s*(\d{4}-\d{2}-\d{2}|\d{1,2}[/\.]\d{1,2}[/\.]\d{4})(?=\s|$|[-–—])', period)
    return parsed_date(match[1]) if match else None


def add_fitted_logo(sheet, path, row, height_points, max_height=None):
    """Inset a proportional image inside column A, including in RTL Excel."""
    logo = Image(path)
    width = int(WIDTHS[0] * 7 + 5)
    height = height_points * 4 / 3
    ratio = min((width - 12) / logo.width, min(height - 8, max_height or height) / logo.height)
    logo.width, logo.height = logo.width * ratio, logo.height * ratio
    logo.anchor = OneCellAnchor(
        _from=AnchorMarker(col=0, row=row - 1,
                           colOff=pixels_to_EMU((width - logo.width) / 2),
                           rowOff=pixels_to_EMU((height - logo.height) / 2)),
        ext=XDRPositiveSize2D(pixels_to_EMU(logo.width), pixels_to_EMU(logo.height)),
    )
    sheet.add_image(logo)


def put(sheet, row, col, value):
    if isinstance(value, str) and len(value) > 32767:
        raise ValueError('Text exceeds Excel cell limit')
    cell = sheet.cell(row, col, value)
    if isinstance(value, str):
        cell.data_type = 's'
        cell.number_format = '@'
    return cell


@lru_cache(maxsize=256)
def insurer_logo(name):
    directory = Path(__file__).parent / 'assets' / 'insurers'
    catalog = json.loads((directory / 'catalog.json').read_text(encoding='utf-8'))
    def clean(value):
        return re.sub('[^a-z0-9א-ת]', '', unicodedata.normalize('NFKC', value).lower())
    value = clean(name)
    suffix = r'(?:(?:חברה|חברת)?ל?ביטוח(?:משכנתאות)?|ופיננסים|ופיננסיים|ישראל|בעמ|אגודהשיתופיתמרכזית|insurance|company|limited|ltd)*'
    aliases = sorted(((clean(alias), brand['file']) for brand in catalog for alias in brand['aliases']), key=lambda pair: -len(pair[0]))
    for alias, file in aliases:
        if value.startswith(alias) and re.fullmatch(suffix, value[len(alias):]):
            return directory / file
    return None


def create_export(request):
    book = Workbook()
    sheet = book.active
    sheet.title = 'תיק ביטוח'
    calc = book.create_sheet('נתוני חישוב')
    audit = book.create_sheet('מקור ושינויים')
    calc.sheet_state = audit.sheet_state = 'hidden'
    book.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
    cached = {}
    for tab in book:
        tab.sheet_view.rightToLeft = True
        tab.sheet_view.showGridLines = False
    for col, width in enumerate(WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width

    def style(row, start=1, end=13, fill=None, bold=False):
        for col in range(start, end + 1):
            cell = sheet.cell(row, col)
            cell.font = Font(name='Arial', size=10, color='FFFFFF' if fill in (NAVY, TEAL) else NAVY, bold=bold)
            cell.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True, readingOrder=2)
            cell.fill = PatternFill('solid', fgColor=fill or 'FFFFFF')
            cell.border = Border(bottom=Side(style='hair', color=LINE))

    def band(text, fill=NAVY):
        row = sheet.max_row + 2
        sheet.row_dimensions[row - 1].height = 7
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=13)
        put(sheet, row, 1, text)
        style(row, fill=fill, bold=True)
        sheet.row_dimensions[row].height = 28 if fill == NAVY else 23
        sheet.cell(row, 1).font = Font(name='Arial', size=14 if fill == NAVY else 12, color='FFFFFF' if fill in (NAVY, TEAL) else NAVY, bold=True)
        if text in ('מצב קיים', 'מצב חדש'):
            sheet.cell(row, 1).alignment = Alignment(horizontal='center', vertical='center', readingOrder=2)
        return row

    def logical(row, values, fill=None, bold=False):
        for (left, right), value in zip(SPANS, values):
            if left != right:
                sheet.merge_cells(start_row=row, start_column=left, end_row=row, end_column=right)
            put(sheet, row, left, value)
        style(row, fill=fill, bold=bold)
        for col in (2, 3, 6, 7, 9, 10):
            sheet.cell(row, col).alignment = Alignment(horizontal='center', vertical='center', wrap_text=True, readingOrder=2)

    # Brand occupies its own white block; agent identity and date have room to fit.
    sheet.merge_cells('A1:A2')
    for area in ('B1:C1', 'B2:C2', 'D1:E1', 'D2:E2'):
        sheet.merge_cells(area)
    put(sheet, 1, 2, 'פרטי סוכן')
    put(sheet, 2, 2, ' '.join(filter(None, [request.agent.firstName, request.agent.lastName])))
    put(sheet, 1, 4, 'תאריך')
    put(sheet, 2, 4, parsed_date(request.agent.date) or request.agent.date)
    sheet['D2'].number_format = 'dd/mm/yyyy'
    style(1, start=2, end=5, fill=NAVY, bold=True)
    style(2, start=2, end=5)
    sheet['D2'].number_format = 'dd/mm/yyyy'
    sheet.row_dimensions[1].height = sheet.row_dimensions[2].height = 29
    add_fitted_logo(sheet, Path(__file__).parent / 'assets' / 'bmc-select.jpg', 1, 58)
    band('לקוחות')
    row = sheet.max_row + 1
    customer_headers = ['מספר לקוח בייצוא', 'שם פרטי', 'שם משפחה', 'תעודת זהות', 'קרבה משפחתית',
                        'גיל ביטוחי', 'תאריך לידה', 'מין', 'מצב משפחתי', '', 'עישון', 'תאריך הדוח', 'הערות ייבוא']
    for col, value in enumerate(customer_headers, 1):
        put(sheet, row, col, value)
    sheet.merge_cells(start_row=row, start_column=9, end_row=row, end_column=10)
    style(row, fill=TEAL, bold=True)
    sheet.row_dimensions[row].height = 34
    for col in range(1, 14):
        sheet.cell(row, col).alignment = Alignment(horizontal='center', vertical='center', wrap_text=True, readingOrder=2)
    age_date = parsed_date(request.agent.date) or date.today()
    for number, customer in enumerate(request.customers, 1):
        row += 1
        details = customer.details
        values = [number, details['firstName'], details['lastName'], details['identity'], details['relationship'],
                  insurance_age(details['birthDate'], age_date), parsed_date(details['birthDate']) or details['birthDate'],
                  details['gender'], details['maritalStatus'], '', details['smoking'],
                  (parsed_date(customer.report.report_date or '') or customer.report.report_date or '') if customer.report else '',
                  '\n'.join(customer.report.warnings) if customer.report else '']
        for col, value in enumerate(values, 1):
            put(sheet, row, col, value)
        sheet.merge_cells(start_row=row, start_column=9, end_row=row, end_column=10)
        style(row, fill=PALE if number % 2 else None)
        for col in range(1, 13):
            sheet.cell(row, col).alignment = Alignment(horizontal='center', vertical='center', wrap_text=True, readingOrder=2)
        for col in (7, 12):
            sheet.cell(row, col).number_format = 'dd/mm/yyyy'
        sheet.row_dimensions[row].height = min(409, max(28, max(ceil(len(str(v)) / 17) for v in values) * 14))

    calc.append(['קבוצה', 'מבוטחים', 'חודשי', 'שנתי', 'חישוב חלקי', 'מצב', 'חודשי לפי הבחירה'])
    put(calc, 1, 16, 'אפשרויות כיסוי')
    coverage_option_row = 2
    audit_headers = ['מצב', 'מספר לקוח', 'מזהה כיסוי', 'גיליון מקור', 'שורת מקור', 'מוחרג']
    for header in HEADERS:
        audit_headers.extend([header + ' — מקור', header + ' — מצב קיים', header + ' — מצב זה'])
    audit.append(audit_headers + ['הערות בדיקה'])
    names = {c.id: ' '.join(filter(None, [c.details['firstName'], c.details['lastName']])) or f'לקוח {n}'
             for n, c in enumerate(request.customers, 1)}
    counts = Counter(name.casefold() for name in names.values())
    names = {c.id: names[c.id] + (f' (לקוח {n})' if counts[names[c.id].casefold()] > 1 or names[c.id] == FAMILY else '')
             for n, c in enumerate(request.customers, 1)}

    def formula(cell, expression, value):
        cell.value = expression
        cell.number_format = MONEY
        cached[cell.coordinate] = value

    def chunks(text, width, max_lines=18):
        # Keep every character, using continuation rows before Excel's height limit.
        result, current, lines = [], '', 0
        for part in text.splitlines(keepends=True) or ['']:
            while part:
                piece, part = part[:width], part[width:]
                if lines >= max_lines:
                    result.append(current)
                    current, lines = '', 0
                current += piece
                lines += 1
        result.append(current)
        return result

    group_number = 0
    audit_row = 2
    customer_totals = {}
    for state in ['מצב קיים', 'מצב חדש']:
        band(state)
        if state == 'מצב חדש':
            r = sheet.max_row + 1
            sheet.merge_cells(start_row=r, start_column=1, end_row=r, end_column=13)
            put(sheet, r, 1, 'הצעה בלבד. נכללים רק כיסויים שהועתקו; כיסוי שלא הועתק אינו נחשב כמבוטל.')
            style(r)
            sheet.row_dimensions[r].height = 24
        grouped = {}
        for owner, customer in enumerate(request.customers, 1):
            for entry in customer.report.entries if customer.report else []:
                if state == 'מצב חדש' and not entry.copied:
                    continue
                baseline = entry.baseline or (entry.original if entry.copied else entry.values)
                values = baseline if state == 'מצב קיים' else entry.values
                excluded = entry.excluded if state == 'מצב חדש' or not entry.copied else False
                normalized, problems = normalize(values)
                category = values['category'].strip()
                section = 'ביטוח אלמנטרי' if re.search('רכב|דירה|דירות|עסק|רכוש|אחריות|צד ג|אלמנטרי', category) else 'ביטוחי פרט'
                complete_key = all(values[k].strip() for k in ['insured_id', 'insurer', 'policy_number'])
                key = (section, values['insurer'].strip(), values['policy_number'].strip())
                if not complete_key:
                    key += (customer.id, entry.id)
                grouped.setdefault(key, []).append((customer, entry, values, normalized, problems, excluded))
                audit_values = [state, owner, entry.id, entry.source_sheet, entry.source_row, 'כן' if excluded else 'לא']
                for field in FIELDS:
                    audit_values.extend([entry.original[field], baseline[field], values[field]])
                audit_values.append('\n'.join(entry.issues + problems))
                for col, value in enumerate(audit_values, 1):
                    put(audit, audit_row, col, value)
                audit_row += 1
        family_rows, selection_rows, annual_rows = [], [], []
        state_partial = False
        for section in ['ביטוחי פרט', 'ביטוח אלמנטרי']:
            matches = [(key, rows) for key, rows in grouped.items() if key[0] == section]
            if not matches and (grouped or section == 'ביטוח אלמנטרי'):
                continue
            band(section, TEAL)
            logical(sheet.max_row + 1, TITLES, fill=NAVY, bold=True)
            sheet.row_dimensions[sheet.max_row].height = 30
            if not matches:
                logical(sheet.max_row + 1, ['', '', '', 'אין כיסויים', '', '', '', '', ''])
            for key, entries in matches:
                group_number += 1
                people = {}
                coverage_options, note_lines = [], []
                for customer, entry, values, normalized, problems, excluded in entries:
                    amounts = people.setdefault(customer.id, [Decimal(0), Decimal(0), False])
                    valid = not any('הפרמיה' in p for p in problems)
                    freq = normalized['frequency']
                    if not excluded:
                        if valid and freq in ('monthly', 'annual'):
                            amounts[0 if freq == 'monthly' else 1] += Decimal(normalized['premium'])
                        else:
                            amounts[2] = True
                    # Short coverage names form a browsing list; all owners, dates,
                    # categories and duplicate source rows remain in the audit sheet.
                    label = (values['product_type'].strip() or values['subcategory'].strip()
                             or values['category'].strip() or 'כיסוי ללא תיאור')
                    cancelled = state == 'מצב חדש' and excluded
                    if excluded:
                        label += ' (מבוטל)' if cancelled else ' (מוחרג)'
                    coverage_options.append(label)
                    hide_zero_annual = (section == 'ביטוח אלמנטרי' and valid
                                        and freq == 'annual' and Decimal(normalized['premium']) == 0)
                    if cancelled:
                        note_lines.append(f'{names[customer.id]}: ביטול פוליסה')
                    details = values['additional_details']
                    if hide_zero_annual:
                        # Hide only this standalone source annotation in the presentation.
                        # The full entry and its notes remain in the audit and calculations.
                        details = '\n'.join(line for line in details.splitlines()
                                            if not re.fullmatch(r'\s*\*\s*פוליסת\s+נהג\s+נקוב\s*', line))
                    if details:
                        note_lines.append(f'{names[customer.id]}: {details}')
                    if not valid or freq not in ('monthly', 'annual'):
                        note_lines.append(f'{names[customer.id]}: פרמיה או תדירות לא תקינה — {values["premium"]} / {values["frequency"]}')
                    if valid and freq == 'annual' and not hide_zero_annual:
                        note_lines.append(f'{names[customer.id]}: פרמיה שנתית {normalized["premium"]} ₪' + (' (מבוטלת)' if cancelled else ' (מוחרגת)' if excluded else ''))
                for cid, amounts in people.items():
                    owner_totals = customer_totals.setdefault((state, cid), [Decimal(0), Decimal(0), False])
                    owner_totals[0] += amounts[0]
                    owner_totals[1] += amounts[1]
                    owner_totals[2] |= amounts[2]
                family = [sum((a[i] for a in people.values()), Decimal(0)) for i in (0, 1)]
                partial = any(a[2] for a in people.values())
                state_partial |= partial
                if partial:
                    note_lines.append('הסכומים חלקיים — נדרש בירור')
                start = calc.max_row + 1
                for label, amounts in [(FAMILY, [*family, partial]), *((names[cid], amount) for cid, amount in people.items())]:
                    cr = calc.max_row + 1
                    for col, value in enumerate([group_number, label, amounts[0], amounts[1], int(amounts[2]), state], 1):
                        put(calc, cr, col, value)
                end = calc.max_row
                family_rows.append(start)
                annual_rows.append(start)
                name = f'Insured_{group_number}'
                book.defined_names.add(DefinedName(name, attr_text=f"'נתוני חישוב'!$B${start}:$B${end}"))
                choice_id = request.selected_customer_id if request.selected_customer_id in people else None
                choice = names[choice_id] if choice_id else FAMILY
                r = sheet.max_row + 1
                categories = list(dict.fromkeys(v['category'].strip() for _, _, v, _, _, _ in entries if v['category'].strip()))
                coverage_title = categories[0] if len(categories) == 1 and len(categories[0]) <= 28 else f'כל הכיסויים ({len(entries)})'
                # Deduplicate display labels only, never the source rows or premiums.
                options = list(dict.fromkeys([coverage_title, *coverage_options]))
                option_start = coverage_option_row
                for option in options:
                    put(calc, coverage_option_row, 16, option)
                    coverage_option_row += 1
                coverage_name = f'Coverages_{group_number}'
                book.defined_names.add(DefinedName(coverage_name, attr_text=f"'נתוני חישוב'!$P${option_start}:$P${coverage_option_row - 1}"))
                coverage_validation = DataValidation(type='list', formula1=f'={coverage_name}', allow_blank=False)
                coverage_validation.showDropDown = False  # False means Excel displays its dropdown arrow.
                coverage_validation.showErrorMessage = True
                coverage_validation.errorStyle = 'stop'
                coverage_validation.errorTitle = 'בחירת כיסוי'
                coverage_validation.error = 'יש לבחור כיסוי מהרשימה.'
                coverage_validation.showInputMessage = True
                coverage_validation.promptTitle = 'כיסויי הפוליסה'
                coverage_validation.prompt = 'הרשימה מציגה את הכיסויים בפוליסה. בחירת כיסוי אינה משנה את הפרמיה או את הסכומים.'
                sheet.add_data_validation(coverage_validation)
                coverage_validation.add(sheet.cell(r, 4))
                periods = list(dict.fromkeys(v['period'].strip() for _, _, v, _, _, _ in entries))
                starts = list(dict.fromkeys(start_date(p).strftime('%d/%m/%Y') if start_date(p) else p or 'לא צוין' for p in periods))
                classifications = list(dict.fromkeys(v['classification'].strip() or 'לא צוין' for _, _, v, _, _, _ in entries))
                detail_chunks = list(zip_longest(chunks('\n'.join(starts), 16, 14),
                                                chunks('\n'.join(classifications), 18, 14),
                                                chunks('\n'.join(note_lines), 38, 14), fillvalue=''))
                dates_text, classification_text, notes_text = detail_chunks[0]
                logical(r, ['', key[1], key[2], coverage_title, dates_text, choice, None, classification_text, notes_text], fill=PALE if group_number % 2 else None)
                sheet.cell(r, 4).alignment = Alignment(horizontal='right', vertical='center', wrap_text=False, readingOrder=2)
                if len(starts) == 1 and parsed_date(starts[0]):
                    sheet.cell(r, 6).value = parsed_date(starts[0])
                    sheet.cell(r, 6).number_format = 'dd/mm/yyyy'
                def content_height(dates, classifications, notes, insurer=''):
                    lines = max(sum(max(1, ceil(len(line) / width)) for line in text.split('\n'))
                                for text, width in [(dates, 16), (classifications, 18), (notes, 38), (insurer, 20)])
                    return max(32, min(350, lines * 16 + 12))
                sheet.row_dimensions[r].height = content_height(dates_text, classification_text, notes_text, key[1])
                logo_path = insurer_logo(key[1])
                if logo_path:
                    add_fitted_logo(sheet, logo_path, r, sheet.row_dimensions[r].height, max_height=38)
                validation = DataValidation(type='list', formula1=f'={name}', allow_blank=False)
                validation.showErrorMessage = True
                validation.errorStyle = 'stop'
                validation.errorTitle = 'בחירת מבוטחים'
                validation.error = 'יש לבחור ערך מהרשימה.'
                validation.showInputMessage = True
                validation.promptTitle = 'פרמיה לפי מבוטחים'
                validation.prompt = 'הבחירה משנה את הפרמיה בשורה. סך המשפחה נשאר ללא שינוי.'
                sheet.add_data_validation(validation)
                validation.add(sheet.cell(r, 7))
                formula(sheet.cell(r, 9), f'=INDEX(\'נתוני חישוב\'!$C${start}:$C${end},MATCH(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(G{r},"~","~~"),"*","~*"),"?","~?"),\'נתוני חישוב\'!$B${start}:$B${end},0))', people[choice_id][0] if choice_id else family[0])
                selection_rows.append(r)
                calc.cell(start, 7, f"='תיק ביטוח'!I{r}")
                for dates_text, classification_text, notes_text in detail_chunks[1:]:
                    rr = sheet.max_row + 1
                    logical(rr, ['', key[1], key[2], '', dates_text, 'המשך פירוט', '', classification_text, notes_text], fill=PALE if group_number % 2 else None)
                    sheet.row_dimensions[rr].height = content_height(dates_text, classification_text, notes_text, key[1])
        for title, rows, column, selected_total in [
            ('סה"כ בחודש למשפחה', family_rows, 'C', False),
            ('סה"כ חודשי לפי הבחירה', selection_rows, 'I', True),
            ('סה"כ פרמיות שנתיות למשפחה (לא תשלום חודשי)', annual_rows, 'D', False),
        ]:
            r = sheet.max_row + 1
            sheet.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
            sheet.merge_cells(start_row=r, start_column=11, end_row=r, end_column=13)
            put(sheet, r, 1, title)
            put(sheet, r, 11, 'סכום חלקי — נדרש בירור' if state_partial else '')
            style(r, fill=BLUE, bold=True)
            sheet.row_dimensions[r].height = 25
            value = sum((cached[f'I{n}'] if selected_total else calc[f'{column}{n}'].value for n in rows), Decimal(0))
            source_column = 'G' if selected_total else column
            end = calc.max_row
            expression = f"=SUMIFS('נתוני חישוב'!{source_column}$2:{source_column}${end},'נתוני חישוב'!B$2:B${end},\"{FAMILY}\",'נתוני חישוב'!F$2:F${end},\"{state}\")"
            formula(sheet.cell(r, 9), expression, value)
            if column == 'D' and value == 0:
                sheet.row_dimensions[r].hidden = True

    for col, label in enumerate(['מצב', 'לקוח', 'תעודת זהות', 'סה״כ חודשי', 'סה״כ שנתי', 'שלמות הסכומים'], 9):
        put(calc, 1, col, label)
    summary_row = 2
    for state in ['מצב קיים', 'מצב חדש']:
        for customer in request.customers:
            monthly, annual, partial = customer_totals.get((state, customer.id), [Decimal(0), Decimal(0), False])
            for col, value in enumerate([state, names[customer.id], customer.details['identity'], monthly, annual, 'חלקי' if partial else 'מלא'], 9):
                put(calc, summary_row, col, value)
            calc.cell(summary_row, 12).number_format = MONEY
            calc.cell(summary_row, 13).number_format = MONEY
            summary_row += 1

    meeting_header = band('סיכום פגישה והמלצות', 'DCEEF3')
    meeting_body = meeting_header + 1
    style(meeting_body, fill='DCEEF3')
    outline = Side(style='thin', color=NAVY)
    for r in (meeting_header, meeting_body):
        for col in range(1, 14):
            sheet.cell(r, col).border = Border(
                top=outline if r == meeting_header else Side(),
                bottom=outline if r == meeting_body else Side(),
                left=outline if col == 1 else Side(),
                right=outline if col == 13 else Side(),
            )
    sheet.merge_cells(start_row=meeting_body, start_column=1, end_row=meeting_body, end_column=13)
    sheet.cell(meeting_body, 1).alignment = Alignment(horizontal='right', vertical='top', wrap_text=True, readingOrder=2)
    sheet.row_dimensions[meeting_body].height = 120
    sheet.sheet_properties.tabColor = NAVY
    sheet.sheet_view.zoomScale = 85
    sheet.freeze_panes = 'A6'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(left=.2, right=.2, top=.25, bottom=.25, header=0, footer=0)
    # Keep the summary heading with its writing box when printing at one-page width.
    # A4 landscape dimensions are in points; Excel column widths use 7-pixel
    # characters plus 5 pixels of padding with the workbook's default Calibri 11.
    printable_width = 297 / 25.4 * 72 - (sheet.page_margins.left + sheet.page_margins.right) * 72
    printable_height = 210 / 25.4 * 72 - (sheet.page_margins.top + sheet.page_margins.bottom) * 72
    width_points = sum(width * 7 + 5 for width in WIDTHS) * .75
    print_scale = min(1, printable_width / width_points)
    # Leave a small allowance for printer/Excel rounding of widths and heights.
    page_height = printable_height / print_scale - 15
    used_height = 0
    for r in range(1, meeting_header):
        dimension = sheet.row_dimensions[r]
        height = 0 if dimension.hidden else dimension.height or sheet.sheet_format.defaultRowHeight
        if used_height + height > page_height:
            used_height = 0
        used_height += height
    if used_height + sheet.row_dimensions[meeting_header].height + sheet.row_dimensions[meeting_body].height > page_height:
        sheet.row_breaks.append(Break(id=meeting_header - 1))
    sheet.print_area = f'A1:M{sheet.max_row}'
    output = BytesIO()
    book.save(output)
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    result = BytesIO()
    with ZipFile(output) as source, ZipFile(result, 'w', ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == 'xl/worksheets/sheet1.xml':
                root = ET.fromstring(data)
                for cell in root.findall('.//s:c', ns):
                    if cell.attrib['r'] in cached:
                        cell.find('s:v', ns).text = str(cached[cell.attrib['r']])
                data = ET.tostring(root, encoding='utf-8')
            target.writestr(item, data)
    return result.getvalue()
