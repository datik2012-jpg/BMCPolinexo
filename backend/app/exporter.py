"""Generate portfolio snapshots entirely in memory; never interpret user text as formulas."""
from collections import Counter
from decimal import Decimal
from io import BytesIO
from math import ceil
import re
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field, field_validator, model_validator

from .parser import FIELDS, HEADERS, normalize


class ExportEntry(BaseModel):
    id: str = Field(max_length=200)
    source_sheet: str = Field(max_length=32767)
    source_row: int = Field(ge=1, le=10000)
    original: dict[str, str]
    values: dict[str, str]
    issues: list[str] = Field(default_factory=list, max_length=100)
    excluded: bool = False

    @field_validator('original', 'values')
    @classmethod
    def complete_fields(cls, value):
        if set(value) != set(FIELDS):
            raise ValueError('Invalid fields')
        return value


class ExportReport(BaseModel):
    instance_id: str
    report_date: str | None
    entries: list[ExportEntry] = Field(min_length=1, max_length=10000)
    warnings: list[str] = Field(default_factory=list, max_length=1000)


DETAILS = {'firstName': 'שם פרטי', 'lastName': 'שם משפחה', 'identity': 'תעודת זהות',
           'relationship': 'קרבה משפחתית', 'birthDate': 'תאריך לידה', 'gender': 'מין',
           'maritalStatus': 'מצב משפחתי', 'smoking': 'עישון'}


class ExportCustomer(BaseModel):
    id: str = Field(max_length=200)
    details: dict[str, str]
    report: ExportReport

    @field_validator('details')
    @classmethod
    def complete_details(cls, value):
        if set(value) != set(DETAILS):
            raise ValueError('Invalid details')
        return value


class ExportRequest(BaseModel):
    customers: list[ExportCustomer] = Field(min_length=1, max_length=50)

    @model_validator(mode='after')
    def bounded(self):
        if len({c.id for c in self.customers}) != len(self.customers):
            raise ValueError('Duplicate customer')
        if sum(len(c.report.entries) for c in self.customers) > 20000:
            raise ValueError('Too many entries')
        # Excel truncates long text silently; reject instead of losing notes.
        def validate(value):
            if isinstance(value, str) and (len(value) > 32767 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]', value)):
                raise ValueError('Unsupported text')
            if isinstance(value, dict):
                for item in value.values():
                    validate(item)
            if isinstance(value, list):
                for item in value:
                    validate(item)
        validate(self.model_dump())
        return self


def append_text_safe(sheet, values):
    if any(isinstance(value, str) and len(value) > 32767 for value in values):
        raise ValueError('Text exceeds Excel cell limit')
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'
            cell.number_format = '@'


def create_export(request: ExportRequest) -> bytes:
    book = Workbook()
    book.remove(book.active)
    customers = book.create_sheet('לקוחות')
    current = book.create_sheet('כיסויים מעודכנים')
    original = book.create_sheet('מקור ושינויים')
    summary = book.create_sheet('סיכומים')
    append_text_safe(customers, ['מספר לקוח בייצוא', *DETAILS.values(), 'תאריך הדוח', 'הערות ייבוא'])
    append_text_safe(current, ['מספר לקוח בייצוא', 'שם פרטי', 'שם משפחה', 'תעודת זהות לקוח',
                                *HEADERS, 'מצב', 'שדות שתוקנו', 'בדיקה', 'גיליון מקור', 'שורת מקור',
                                'פרמיה חודשית לחישוב (₪)', 'פרמיה שנתית לחישוב (₪)'])
    audit_headers = ['מספר לקוח בייצוא', 'שם פרטי', 'שם משפחה', 'מזהה שורה', 'גיליון מקור', 'שורת מקור']
    for header in HEADERS:
        audit_headers.extend([f'{header} — מקור', f'{header} — מעודכן'])
    append_text_safe(original, [*audit_headers, 'מצב', 'הערות המקור'])
    append_text_safe(summary, ['מספר לקוח בייצוא', 'שם פרטי', 'שם משפחה', 'סה״כ חודשי (₪)', 'סה״כ שנתי (₪)', 'שלמות הסכומים'])
    cached = {}
    for index, customer in enumerate(request.customers, 1):
        d, report = customer.details, customer.report
        append_text_safe(customers, [index, *(d[k] for k in DETAILS), report.report_date or '', '\n'.join(report.warnings)])
        monthly, annual = Decimal(0), Decimal(0)
        incomplete = False
        signatures = Counter(tuple(e.values[k] for k in FIELDS) for e in report.entries if not e.excluded)
        for entry in report.entries:
            normalized, problems = normalize(entry.values)
            valid_amount = not any('הפרמיה' in problem for problem in problems)
            freq = normalized['frequency']
            month = Decimal(normalized['premium']) if valid_amount and freq == 'monthly' else None
            year = Decimal(normalized['premium']) if valid_amount and freq == 'annual' else None
            if not entry.excluded:
                monthly += month or Decimal(0)
                annual += year or Decimal(0)
                incomplete |= not valid_amount or freq not in ('monthly', 'annual')
            status = 'מוחרג' if entry.excluded else 'נכלל'
            changed = ' · '.join(header for key, header in zip(FIELDS, HEADERS) if entry.values[key] != entry.original[key])
            if signatures[tuple(entry.values[k] for k in FIELDS)] > 1:
                problems.append('חשד לשורה כפולה — נכללת בחישוב עד להחרגה')
            problems.extend(i for i in entry.issues if 'אפסים מובילים' in i or 'נוסח' in i)
            append_text_safe(current, [index, d['firstName'], d['lastName'], d['identity'],
                                      *(entry.values[k] for k in FIELDS), status, changed, '\n'.join(problems),
                                      entry.source_sheet, entry.source_row, month, year])
            values = [index, d['firstName'], d['lastName'], entry.id, entry.source_sheet, entry.source_row]
            for key in FIELDS:
                values.extend([entry.original[key], entry.values[key]])
            append_text_safe(original, [*values, status, '\n'.join(entry.issues)])
        append_text_safe(summary, [index, d['firstName'], d['lastName'], None, None,
                                   'חלקי — יש פרמיות או תדירויות לא תקינות' if incomplete else 'מלא'])
        row = summary.max_row
        for col, amount, source_col in [('D', monthly, 'U'), ('E', annual, 'V')]:
            # Match stable customer numbers so sorting coverage rows remains safe.
            end = 1 + sum(len(c.report.entries) for c in request.customers)
            summary[f'{col}{row}'] = f'=SUMIFS(\'כיסויים מעודכנים\'!{source_col}$2:{source_col}${end},\'כיסויים מעודכנים\'!A$2:A${end},A{row},\'כיסויים מעודכנים\'!P$2:P${end},"נכלל")'
            cached[f'{col}{row}'] = amount
    summary.append([])
    append_text_safe(summary, ['סכומים לפי התיקונים וההחרגות בזמן הייצוא. הסכומים השנתיים אינם תשלומים חודשיים.'])
    append_text_safe(summary, ['כל הכיסויים מיוצאים, גם אם הוסתרו בסינון. כפילויות אינן נמחקות אוטומטית.'])
    append_text_safe(summary, ['זהו צילום מצב של העבודה; שינוי פרטי כיסוי בקובץ אינו מעדכן את האפליקציה.'])
    for sheet in book:
        sheet.sheet_view.rightToLeft = True
        sheet.freeze_panes = 'E2' if sheet == current else 'A2'
        last_data_row = len(request.customers) + 1 if sheet == summary else sheet.max_row
        sheet.auto_filter.ref = f'A1:{get_column_letter(sheet.max_column)}{last_data_row}'
        for row in sheet:
            for cell in row:
                cell.font = Font(name='Arial', size=11, color='183D42')
                cell.alignment = Alignment(vertical='top', horizontal='right', wrap_text=True)
                if cell.row == 1:
                    cell.fill = PatternFill('solid', fgColor='216B60')
                    cell.font = Font(name='Arial', size=11, color='FFFFFF', bold=True)
                elif cell.row % 2 == 0:
                    cell.fill = PatternFill('solid', fgColor='EFF6F5')
            sheet.row_dimensions[row[0].row].height = 42
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 24
    current.column_dimensions['K'].width = 60  # Notes stay alongside their coverage.
    original.column_dimensions['S'].width = original.column_dimensions['T'].width = 60
    customers.column_dimensions['K'].width = 60
    summary.column_dimensions['F'].width = 50
    for row in range(len(request.customers) + 3, summary.max_row + 1):
        summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    for sheet in book:
        for row in sheet:
            lines = max((sum(max(1, ceil(len(line) / max(1, sheet.column_dimensions[cell.column_letter].width - 2)))
                             for line in str(cell.value or '').split('\n')) for cell in row if cell.value is not None), default=1)
            sheet.row_dimensions[row[0].row].height = min(409, max(42, lines * 16))
    for sheet, columns, end in [(current, ['U', 'V'], current.max_row), (summary, ['D', 'E'], len(request.customers) + 1)]:
        for col in columns:
            for row in range(2, end + 1):
                sheet[f'{col}{row}'].number_format = '#,##0.00 "₪"'
    output = BytesIO()
    book.save(output)
    # Cache Decimal-computed formula results for previewers; Excel recalculates on open.
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    result = BytesIO()
    with ZipFile(output) as source, ZipFile(result, 'w', ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == 'xl/worksheets/sheet4.xml':
                root = ET.fromstring(data)
                for cell in root.findall('.//s:c', ns):
                    if cell.attrib['r'] in cached:
                        cell.find('s:v', ns).text = str(cached[cell.attrib['r']])
                data = ET.tostring(root, encoding='utf-8')
            target.writestr(item, data)
    return result.getvalue()
