from decimal import Decimal
from io import BytesIO
import re
from zipfile import ZipFile, ZIP_DEFLATED

from fastapi.testclient import TestClient
from openpyxl import Workbook
import pytest

from app.main import app
from app.parser import HEADERS, ImportProblem, parse_workbook

def fixture(rows=None, broken_dimension=True, second_sheet=False):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Synthetic'
    sheet.cell(2, 2, 'דוח בדיקה')
    sheet.cell(2, 6, '01/01/2026')
    sheet.append([])
    for col, header in enumerate(HEADERS, 1):
        sheet.cell(4, col, header)
    if rows is None:
        rows = []
        premiums = [[1587,4055,0,0,0,0], [69.91], [71.21,76.85,19.14,36.68,11.40,14.87,15.98], [60.64], [103.35], [36.86,52.82], [51.05], [142.56]]
        for policy, amounts in enumerate(premiums, 1):
            for index, amount in enumerate(amounts):
                rows.append([999999999, 'קטגוריה אחרת' if policy == 3 and index == 1 else 'קטגוריה לבדיקה', 'משני', 'כיסוי חוזר' if policy == 6 else f'כיסוי {index}', f'מבטח {policy % 2}', f'01/01/202{index % 2 + 5} - 31/12/2030', '', amount, 'שנתית' if policy == 1 else 'חודשית', 800000 + policy, 'אישי'])
    sheet.append([None, 'תחום - כללי'])
    for row in rows:
        sheet.append(row)
    if second_sheet:
        extra = book.create_sheet('Second')
        extra.append(list(HEADERS))
        extra.append(rows[0])
    buffer = BytesIO()
    book.save(buffer)
    if not broken_dimension:
        return buffer.getvalue()
    output = BytesIO()
    with ZipFile(buffer) as original, ZipFile(output, 'w', ZIP_DEFLATED) as modified:
        for item in original.infolist():
            value = original.read(item.filename)
            if item.filename == 'xl/worksheets/sheet1.xml':
                value = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A2:K5"', value)
            modified.writestr(item, value)
    return output.getvalue()

def test_structural_fixture_and_decimal_totals():
    result = parse_workbook(fixture())
    entries = result['entries']
    assert len(entries) == 20
    assert len({tuple(e['values'][k] for k in ('insured_id','insurer','policy_number')) for e in entries}) == 8
    totals = {frequency: sum((Decimal(e['values']['premium']) for e in entries if e['values']['frequency'] == frequency), Decimal(0)) for frequency in ('monthly','annual')}
    assert totals == {'monthly': Decimal('763.32'), 'annual': Decimal('5642')}
    assert sum(e['values']['premium'] == '0' for e in entries) == 4
    assert entries[-1]['source_row'] > 5
    assert result['report_date'] == '01/01/2026'
    assert entries[0]['original']['frequency'] == 'שנתית'
    repeated = [e for e in entries if e['values']['policy_number'] == '800006']
    assert len(repeated) == 2
    assert repeated[0]['values']['product_type'] == repeated[1]['values']['product_type']
    assert repeated[0]['values']['period'] != repeated[1]['values']['period']
    assert len({e['values']['category'] for e in entries if e['values']['policy_number'] == '800003'}) == 2

def row(**changes):
    from app.parser import FIELDS
    values = dict(zip(FIELDS, [999999999,'קטגוריה','משני','כיסוי','מבטח','מתחדש','',10,'חודשית',123,'אישי']))
    values.update(changes)
    return list(values.values())

def test_duplicates_missing_invalid_preserved():
    result = parse_workbook(fixture([row(), row(), row(policy_number=None, premium=None), row(premium='bad',frequency='?')]))
    assert len(result['entries']) == 4
    assert 'כפולה' in result['entries'][0]['issues'][-1]
    assert any('שיוך' in issue for issue in result['entries'][2]['issues'])
    assert result['entries'][3]['values']['premium'] == 'bad'
    assert any('תדירות' in issue for issue in result['entries'][3]['issues'])
    assert any('תקינה' in issue for issue in result['entries'][3]['issues'])

@pytest.mark.parametrize('premium', ['1,2', '-1', 'NaN', 'Infinity', '1.234', '1000000000000'])
def test_invalid_decimal_formats(premium):
    result = parse_workbook(fixture([row(premium=premium)]))
    assert any('תקינה' in issue for issue in result['entries'][0]['issues'])

@pytest.mark.parametrize('rows', [[row(),row(insured_id=888888888)], [row(premium='=1+1')]])
def test_rejected_data(rows):
    with pytest.raises(ImportProblem):
        parse_workbook(fixture(rows))

def test_multiple_data_sheets_rejected():
    with pytest.raises(ImportProblem):
        parse_workbook(fixture(second_sheet=True))

def test_api_and_safe_errors():
    client = TestClient(app)
    instance = client.get('/api/instance')
    assert instance.headers['cache-control'] == 'no-store'
    response = client.post('/api/import', files={'file':('synthetic.xlsx', fixture())})
    assert response.status_code == 200
    assert response.json()['instance_id'] == instance.json()['instance_id']
    assert response.headers['cache-control'] == 'no-store'
    for filename, payload in [('secret.xlsx',b'sensitive contents'),('secret.csv',b'x')]:
        error = client.post('/api/import',files={'file':(filename,payload)})
        assert error.status_code == 400
        assert 'secret' not in error.text and 'sensitive' not in error.text
    assert client.post('/api/import', files=[('file',('a.xlsx',fixture())),('file',('b.xlsx',fixture()))]).status_code == 400
    assert client.post('/api/import', files={'file':('large.xlsx',b'x' * (5*1024*1024))}).status_code == 413

def test_archive_limits(monkeypatch):
    import app.parser as parser
    monkeypatch.setattr(parser,'MAX_EXPANDED',10)
    with pytest.raises(ImportProblem):
        parse_workbook(fixture())

def test_row_limit(monkeypatch):
    import app.parser as parser
    monkeypatch.setattr(parser,'MAX_ROWS',6)
    with pytest.raises(ImportProblem):
        parse_workbook(fixture())

def test_no_spooling_or_sensitive_logs(monkeypatch, caplog):
    import tempfile
    def forbidden(*args, **kwargs):
        raise AssertionError('Upload must stay in memory')
    payload = fixture([row(additional_details='PRIVATE_CONTENT_SENTINEL')])
    monkeypatch.setattr(tempfile, 'TemporaryFile', forbidden)
    monkeypatch.setattr(tempfile, 'SpooledTemporaryFile', forbidden)
    client = TestClient(app)
    response = client.post('/api/import',files={'file':('PRIVATE_FILENAME.xlsx',payload)})
    assert response.status_code == 200
    assert all(value not in caplog.text for value in ('PRIVATE_CONTENT_SENTINEL','PRIVATE_FILENAME','999999999'))

def test_xml_entities_rejected():
    output = BytesIO()
    with ZipFile(BytesIO(fixture())) as source, ZipFile(output,'w',ZIP_DEFLATED) as target:
        for item in source.infolist():
            value = source.read(item.filename)
            if item.filename == 'xl/worksheets/sheet1.xml':
                value = b'<!DOCTYPE worksheet [<!ENTITY xxe SYSTEM "file:///private">]>' + value
            target.writestr(item,value)
    with pytest.raises(ImportProblem):
        parse_workbook(output.getvalue())

def test_numeric_identifier_padding():
    book = Workbook()
    sheet = book.active
    sheet.append(list(HEADERS))
    sheet.append(row(insured_id=123, policy_number=456))
    sheet['A2'].number_format = '000000000'
    sheet['J2'].number_format = '000000'
    buffer = BytesIO()
    book.save(buffer)
    entry = parse_workbook(buffer.getvalue())['entries'][0]
    assert entry['original']['insured_id'] == '000000123'
    assert entry['values']['policy_number'] == '000456'
    assert not entry['issues']
