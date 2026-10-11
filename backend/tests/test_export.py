from copy import deepcopy
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app, INSTANCE_ID
from app.parser import parse_workbook
from test_import import fixture, row


def customer(identity='000000001', first='בדיקה'):
    report = parse_workbook(fixture([row(insured_id=identity, premium='10.10'),
                                     row(insured_id=identity, premium='120', frequency='שנתית'),
                                     row(insured_id=identity, premium='0')]))
    report['instance_id'] = INSTANCE_ID
    return {'id': identity, 'details': {'firstName': first, 'lastName': 'סינתטי', 'identity': identity,
            'relationship': '', 'birthDate': '', 'gender': 'זכר', 'maritalStatus': '', 'smoking': ''}, 'report': report}


def test_export_rejects_bad_or_stale_payload_without_echoing_data():
    client = TestClient(app)
    for body in [{}, {'customers': []}, {'customers': [customer(), customer()]}]:
        assert client.post('/api/export', json=body).status_code == 400
    c = customer()
    c['report']['instance_id'] = 'old'
    assert client.post('/api/export', json={'customers': [c]}).status_code == 409
    c['report']['instance_id'] = INSTANCE_ID
    c['report']['entries'][0]['values']['additional_details'] = 'x' * 32768
    result = client.post('/api/export', json={'customers': [c]})
    assert result.status_code == 400
    assert 'xxx' not in result.text


def exported(customers, **kwargs):
    response = TestClient(app).post('/api/export', json={'customers': customers, **kwargs})
    assert response.status_code == 200
    assert response.headers['cache-control'] == 'no-store'
    assert response.headers['x-instance-id'] == INSTANCE_ID
    return load_workbook(BytesIO(response.content)), load_workbook(BytesIO(response.content), data_only=True)


def totals(book, title):
    return [book.active.cell(c.row, 9).value for row in book.active for c in row if c.value == title]


def policy_rows(book):
    return [book.active[str(v.sqref)] for v in book.active.data_validations.dataValidation if v.formula1.startswith('=Insured_')]


def test_states_formulas_owners_and_history():
    a, b = customer(), customer('000000002', 'שני')
    entry = a['report']['entries'][0]
    entry['baseline'] = deepcopy(entry['values'])
    entry['copied'] = True
    entry['values']['additional_details'] = '=HYPERLINK("https://example.com")\nפרטים לכיסוי'
    entry['values']['premium'] = '20.20'
    a['report']['entries'][1]['copied'] = True
    a['report']['entries'][1]['excluded'] = True
    snapshot = deepcopy([a, b])
    book, values = exported([a, b], agent={'firstName': '=Test', 'lastName': 'Agent', 'date': '2026-10-04'}, selected_customer_id=a['id'])
    assert [s.sheet_state for s in book] == ['visible', 'hidden', 'hidden']
    sheet = book.active
    assert sheet['B2'].value == '=Test Agent' and sheet['B2'].data_type == 's'
    assert sheet['D6'].value == '000000001' and sheet['D7'].value == '000000002'
    assert sheet.page_setup.fitToWidth == 1 and sheet.page_setup.fitToHeight == 0
    assert sheet.sheet_view.rightToLeft
    assert totals(values, 'סה"כ בחודש למשפחה') == [20.2, 20.2]
    assert totals(values, 'סה"כ חודשי לפי הבחירה') == [10.1, 20.2]
    assert totals(values, 'סה"כ פרמיות שנתיות למשפחה (לא תשלום חודשי)') == [240, 0]
    rows = policy_rows(book)
    assert len(rows) == 2
    assert all(c.value == 'בדיקה סינתטי' for c in rows)
    assert sheet.cell(rows[0].row, 9).data_type == 'f'
    assert 'HYPERLINK' not in (sheet.cell(rows[0].row, 11).value or '')
    assert 'HYPERLINK' in sheet.cell(rows[1].row, 11).value
    assert sheet.cell(rows[1].row, 11).data_type == 's'
    assert [a, b] == snapshot
    assert book['מקור ושינויים'].max_row == 9
    assert sheet._images


def test_empty_proposal_manual_customer_and_invalid_amount():
    a, b = customer(), customer('000000002', 'ריק')
    b['report'] = None
    a['report']['entries'][0]['values']['premium'] = 'bad'
    book, values = exported([a, b])
    assert values.active['B7'].value == 'ריק'
    assert totals(values, 'סה"כ בחודש למשפחה') == [0, 0]
    assert len(policy_rows(book)) == 1
    assert any('bad' in str(c.value) for row in values.active for c in row)
    assert any('חלקי' in str(c.value) for row in values.active for c in row)


def test_incomplete_keys_duplicates_zero_and_policy_separation():
    a = customer()
    entries = []
    for n, policy in enumerate(['001', '001', '002', '', '']):
        e = deepcopy(a['report']['entries'][0])
        e['id'] = str(n)
        e['values']['policy_number'] = policy
        e['values']['premium'] = '0' if n == 2 else '10.10'
        entries.append(e)
    a['report']['entries'] = entries
    book, values = exported([a])
    assert len(policy_rows(book)) == 4
    assert totals(values, 'סה"כ בחודש למשפחה') == [40.4, 0]
    assert book.active.cell(policy_rows(book)[0].row, 3).value == '001'


def test_duplicate_names_and_long_notes():
    a, b = customer(), customer('000000002')
    note = 'טקסט ארוך ' * 2000
    a['report']['entries'][0]['values']['additional_details'] = note
    book, _ = exported([a, b])
    choices = [r[1].value for r in book['נתוני חישוב'].iter_rows(min_row=2)]
    assert 'בדיקה סינתטי (לקוח 1)' in choices and 'בדיקה סינתטי (לקוח 2)' in choices
    visible = ''.join(str(book.active.cell(r, 11).value or '') for r in range(1, book.active.max_row + 1))
    assert note in visible
    assert max(d.height or 0 for d in book.active.row_dimensions.values()) <= 409


def test_general_insurance_owner_totals_and_literal_wildcard_names():
    a, b = customer(first='*'), customer('000000002', '?')
    for c in [a, b]:
        c['report']['entries'][1]['values']['category'] = 'ביטוח רכב'
        c['report']['entries'][1]['values']['policy_number'] = '000car'
    book, values = exported([a, b], selected_customer_id=a['id'])
    rows = policy_rows(book)
    assert len(rows) == 2
    car_row = next(c.row for c in rows if book.active.cell(c.row, 3).value == '000car')
    assert any(book.active.cell(r, 1).value == 'ביטוח אלמנטרי' for r in range(1, car_row))
    assert 'SUBSTITUTE' in book.active.cell(rows[0].row, 9).value
    assert totals(values, 'סה"כ בחודש למשפחה') == [20.2, 0]
    calc = values['נתוני חישוב']
    assert calc['L2'].value == 10.1 and calc['M2'].value == 120
    assert calc['L3'].value == 10.1 and calc['M3'].value == 120
    assert calc['K2'].value == '000000001'


def test_coverage_dropdown_keeps_many_coverages_in_one_compact_policy_row():
    a = customer()
    entries = []
    for n in range(40):
        e = deepcopy(a['report']['entries'][0])
        e['id'] = f'coverage-{n}'
        e['source_row'] = n + 2
        e['values'].update(category='ביטוח בריאות', product_type=f'כיסוי בדיקה {n}',
                           period='תקופה מקורית', premium='1.25')
        e['baseline'] = deepcopy(e['values'])
        e['copied'] = True
        entries.append(e)
    a['report']['entries'] = entries
    book, values = exported([a])
    selectors = [v for v in book.active.data_validations.dataValidation if v.formula1.startswith('=Coverages_')]
    assert len(selectors) == 2  # Exactly one policy row per state, independent of coverage count.
    for validation in selectors:
        cell = book.active[str(validation.sqref)]
        assert cell.column == 4 and cell.value == 'ביטוח בריאות'
        assert validation.type == 'list' and validation.showDropDown is False
        assert book.active.row_dimensions[cell.row].height == 32
        assert not cell.alignment.wrap_text
        source_sheet, source_range = next(book.defined_names[validation.formula1[1:]].destinations)
        choices = [r[0].value for r in book[source_sheet][source_range]]
        assert choices == ['ביטוח בריאות', *(f'כיסוי בדיקה {n}' for n in range(40))]
    assert totals(values, 'סה"כ בחודש למשפחה') == [50, 50]
    assert book['מקור ושינויים'].max_row == 81
    assert any(c.value == 'תקופה מקורית' for row in book['מקור ושינויים'] for c in row)
    assert not any(c.value == 'המשך' for row in book.active for c in row)


def test_coverage_choices_are_policy_and_state_specific_and_never_deduplicate_money():
    a = customer()
    first = a['report']['entries'][0]
    first['values']['product_type'] = 'ניתוחים'
    first['baseline'] = deepcopy(first['values'])
    first['copied'] = True
    first['values']['product_type'] = 'טיפולים חדשים'
    duplicate = deepcopy(first)
    duplicate['id'] = 'duplicate'
    a['report']['entries'] = [first, duplicate]
    book, values = exported([a])
    options = []
    for validation in book.active.data_validations.dataValidation:
        if validation.formula1.startswith('=Coverages_'):
            name, cells = next(book.defined_names[validation.formula1[1:]].destinations)
            options.append([r[0].value for r in book[name][cells]])
    assert 'ניתוחים' in options[0] and 'טיפולים חדשים' not in options[0]
    assert 'טיפולים חדשים' in options[1] and 'ניתוחים' not in options[1]
    assert len(options[0]) == len(options[1]) == 2  # One summary and one unique display label.
    assert totals(values, 'סה"כ בחודש למשפחה') == [20.2, 20.2]
    assert book['מקור ושינויים'].max_row == 5  # Both source coverages in both states.


@pytest.mark.parametrize('note', ['', 'הוראות שמורות\n=1+1'])
@pytest.mark.parametrize('frequency', ['monthly', 'annual'])
def test_cancel_restore_exports_instruction_without_changing_notes_or_other_owners(note, frequency):
    a, b = customer(), customer('000000002', 'שני')
    for owner in [a, b]:
        entry = owner['report']['entries'][0]
        entry['values']['frequency'] = frequency
        entry['baseline'] = deepcopy(entry['values'])
        entry['copied'] = True
        owner['report']['entries'] = [entry]
    entry = a['report']['entries'][0]
    entry['values'].update(premium='25.25', additional_details=note)
    before_cancel = deepcopy(entry)
    entry['excluded'] = True
    book, values = exported([a, b])
    rows = policy_rows(book)
    existing_notes = book.active.cell(rows[0].row, 11).value or ''
    proposed_notes = book.active.cell(rows[1].row, 11)
    assert 'ביטול פוליסה' not in existing_notes
    assert proposed_notes.data_type == 's'
    assert 'בדיקה סינתטי: ביטול פוליסה' in proposed_notes.value
    assert 'שני סינתטי: ביטול פוליסה' not in proposed_notes.value
    if note:
        assert note in proposed_notes.value
    assert any(c.value == 'הוראות הסוכן' for row in book.active for c in row)
    assert ' (מבוטל)' in str([r[15].value for r in book['נתוני חישוב']])
    total_label = 'סה"כ בחודש למשפחה' if frequency == 'monthly' else 'סה"כ פרמיות שנתיות למשפחה (לא תשלום חודשי)'
    assert totals(values, total_label) == [20.2, 10.1]
    # Cancellation is a flag: stored source, baseline and manual notes stay intact.
    audit = book['מקור ושינויים']
    proposed_a = next(row for row in audit.iter_rows(min_row=2, values_only=True)
                      if row[0] == 'מצב חדש' and row[1] == 1)
    assert proposed_a[5] == 'כן'
    assert (proposed_a[24] or '') == entry['original']['additional_details']
    assert (proposed_a[25] or '') == entry['baseline']['additional_details']
    assert (proposed_a[26] or '') == note
    entry['excluded'] = False
    restored, restored_values = exported([a, b])
    restored_notes = restored.active.cell(policy_rows(restored)[1].row, 11).value or ''
    assert 'ביטול פוליסה' not in restored_notes
    if note:
        assert note in restored_notes
    assert totals(restored_values, total_label) == [20.2, 35.35]
    assert {k: v for k, v in entry.items() if k != 'excluded'} == before_cancel


def test_reference_layout_dates_classifications_and_blank_meeting_summary():
    a = customer()
    a['details']['birthDate'] = '1986-10-11'
    first = a['report']['entries'][0]
    first['values'].update(insurer='הראל', period='01/02/2020 - 31/01/2030', classification='אישי')
    second = deepcopy(first)
    second['id'] = 'second-start'
    second['values'].update(period='2021-03-04 - 2031-03-03', classification='קבוצתי')
    a['report']['entries'] = [first, second]
    book, values = exported([a], agent={'firstName': 'סוכן', 'lastName': 'בדיקה', 'date': '2026-10-10'})
    sheet = book.active
    assert sheet.max_column == 13
    assert sheet['B2'].value == 'סוכן בדיקה'
    assert sheet['F6'].value == 39
    assert sheet['G6'].value.strftime('%Y-%m-%d') == '1986-10-11'
    row = policy_rows(book)[0].row
    assert sheet.cell(row, 6).value == '01/02/2020\n04/03/2021'
    assert sheet.cell(row, 10).value == 'אישי\nקבוצתי'
    assert sheet.row_dimensions[row].height >= 44
    assert '01/02/2020 - 31/01/2030' in str(list(book['מקור ושינויים'].values))
    summary = next(c.row for cells in sheet for c in cells if c.value == 'סיכום פגישה והמלצות')
    assert sheet.max_row == summary + 3
    assert not sheet.protection.sheet
    assert all(sheet.cell(r, 1).value is None for r in range(summary + 1, summary + 4))
    assert all(f'A{r}:M{r}' in sheet.merged_cells for r in range(summary + 1, summary + 4))
    assert totals(values, 'סה"כ בחודש למשפחה') == [20.2, 0]
    # Images must fit inside their cells with padding and retain source aspect ratio.
    from PIL import Image as PillowImage
    from app.compact_export import WIDTHS
    for picture in sheet._images:
        anchor = picture.anchor
        assert anchor._from.col == 0 and anchor._from.colOff > 0 and anchor._from.rowOff > 0
        width, height = anchor.ext.cx / 9525, anchor.ext.cy / 9525
        source = PillowImage.open(picture.ref)
        assert width / height == pytest.approx(source.width / source.height, rel=1e-5)
        assert anchor._from.colOff / 9525 + width < WIDTHS[0] * 7 + 5
        available_height = 58 if anchor._from.row == 0 else sheet.row_dimensions[anchor._from.row + 1].height
        assert anchor._from.rowOff / 9525 + height < available_height * 4 / 3


def test_many_dates_and_classifications_preserve_all_values_with_continuations():
    a = customer()
    entries = []
    for n in range(35):
        e = deepcopy(a['report']['entries'][0])
        e['id'] = f'detail-{n}'
        e['values'].update(period=f'01/01/{1990+n} - 31/12/2030', classification=f'סיווג {n}', premium='1')
        entries.append(e)
    a['report']['entries'] = entries
    book, values = exported([a])
    dates = '\n'.join(str(book.active.cell(r, 6).value or '') for r in range(1, book.active.max_row + 1))
    classes = '\n'.join(str(book.active.cell(r, 10).value or '') for r in range(1, book.active.max_row + 1))
    assert all(f'01/01/{1990+n}' in dates and f'סיווג {n}' in classes for n in range(35))
    assert len(policy_rows(book)) == 1
    assert totals(values, 'סה"כ בחודש למשפחה') == [35, 0]
    assert max(d.height or 0 for d in book.active.row_dimensions.values()) < 409


@pytest.mark.parametrize('birth,on_date,expected', [
    ('1986-10-11', '2026-10-10', 39), ('1986-10-11', '2026-10-11', 40),
    ('2024-02-29', '2025-02-28', 0), ('2024-02-29', '2025-03-01', 1),
    ('', '2026-10-10', None), ('bad', '2026-10-10', None), ('2030-01-01', '2026-10-10', None),
])
def test_export_insurance_age_at_last_birthday(birth, on_date, expected):
    from app.compact_export import insurance_age, parsed_date
    assert insurance_age(birth, parsed_date(on_date)) == expected
