"""Verify real browser downloads using synthetic customers and in-memory uploads."""
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect
from browser_check import fixture_module
from customer_helpers import fill_customer

with sync_playwright() as p:
    browser = p.chromium.launch()
    page = browser.new_page(viewport={'width': 1440, 'height': 1000}, accept_downloads=True)
    page.goto('http://127.0.0.1:8080')
    expect(page.get_by_label('שם פרטי', exact=True)).to_be_enabled(timeout=15000)
    for index, identity in enumerate(['999999999', '888888888', '777777777']):
        if index:
            page.get_by_role('button', name='+ הוסף לקוח נוסף', exact=True).click()
        card = page.locator('.customer-card').nth(0 if index == 0 else 1)
        fill_customer(card, identity)
        card.get_by_label('שם פרטי', exact=True).fill(f'בדיקה{index}')
        card.get_by_label('העלאת קובץ Excel הר ביטוח', exact=True).set_input_files({
            'name': 'synthetic.xlsx', 'mimeType': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'buffer': fixture_module.fixture([
                fixture_module.row(insured_id=int(identity), premium=10 + index, category='ביטוח בריאות'),
                fixture_module.row(insured_id=int(identity), premium=120, frequency='שנתית', category='ביטוח רכב'),
            ])})
        expect(card.locator('.upload-success')).to_be_visible()
    section = page.get_by_role('region', name='תיק לקוח בדיקה0 סינתטי', exact=True)
    section.locator('.policy summary').click()
    section.get_by_role('button', name='עריכה', exact=True).click()
    page.get_by_role('dialog').get_by_label('פרטים נוספים', exact=True).fill('פרטים של כיסוי ראשון\n=1+1')
    page.get_by_role('button', name='שמירת תיקונים', exact=True).click()
    section.get_by_role('button', name='החרגה', exact=True).click()

    def download(scope):
        with page.expect_download() as event:
            scope.get_by_role('button', name='ייצוא ל־Excel', exact=True).click()
        item = event.value
        assert item.failure() is None
        data = Path(item.path()).read_bytes()
        item.delete()
        return load_workbook(BytesIO(data), data_only=True)

    book = download(section)
    assert book['לקוחות'].max_row == 2
    assert book['כיסויים מעודכנים'].max_row == 3  # Also the hidden car coverage.
    assert 'פרטים של כיסוי ראשון' in book['כיסויים מעודכנים']['K2'].value
    assert book['כיסויים מעודכנים']['P2'].value == 'מוחרג'
    assert book['סיכומים']['D2'].value == 0
    assert book['סיכומים']['E2'].value == 120
    selection = page.get_by_role('region', name='בחירת לקוחות לתצוגה משותפת')
    selection.get_by_role('checkbox', name='בדיקה0 סינתטי').check()
    selection.get_by_role('checkbox', name='בדיקה1 סינתטי').check()
    selection.get_by_role('button', name='הצגת הלקוחות יחד').click()
    shared = page.locator('.shared-portfolio')
    book = download(shared)
    assert book['לקוחות'].max_row == 3
    assert book['כיסויים מעודכנים'].max_row == 5
    assert {book['לקוחות'].cell(row, 2).value for row in [2, 3]} == {'בדיקה0', 'בדיקה1'}
    assert book['סיכומים']['D3'].value == 11
    for width in [1440, 390, 320]:
        page.set_viewport_size({'width': width, 'height': 1000})
        assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
    page.route('**/api/export', lambda route: route.fulfill(status=500, json={'detail': 'הייצוא נכשל'}))
    shared.get_by_role('button', name='ייצוא ל־Excel', exact=True).click()
    expect(shared.get_by_role('alert')).to_contain_text('הייצוא נכשל')
    assert page.evaluate('localStorage.length + sessionStorage.length') == 0
    browser.close()
print('PASS: single/shared downloads, selected customers only, notes in coverage rows, hidden/excluded rows, annual separation, errors, RTL widths')
